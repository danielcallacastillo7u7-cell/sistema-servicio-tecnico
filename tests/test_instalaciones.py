import unittest
from uuid import uuid4
from io import BytesIO
from openpyxl import load_workbook
import test_multi_equipos as fixtures
from app.models.instalacion import InstalacionCamara

class Instalaciones(unittest.TestCase):
    def setUp(self):
        self.fixture=fixtures.MultiplesEquipos();self.fixture.setUp();self.web=self.fixture.web
        self.data=dict(solicitud_id=str(uuid4()),nombres='Ana',apellidos='Pérez',dni_ruc='12345678',celular='987654321',tecnico_id=1,fecha_instalacion='2026-10-10',direccion='Calle Uno 123, Lima',camaras=[dict(modelo='TP-Link C200',procedencia='Cliente',cantidad=2)],observaciones='Entrada principal',productos=[dict(descripcion='Cable UTP',cantidad='12.50',unidad='Metro')])
    def tearDown(self):
        self.web.close();fixtures.engine.dispose()
    def test_guardado_reporte_independiente_reintento(self):
        r=self.web.post('/camaras/api',json=self.data);self.assertEqual(r.status_code,201,r.text)
        self.assertEqual(self.web.post('/camaras/api',json=self.data).json()['numero'],r.json()['numero'])
        with fixtures.SessionLocal() as db:
            self.assertEqual(db.query(InstalacionCamara).count(),1);self.assertEqual(db.query(fixtures.OrdenServicio).count(),0)
            self.assertIsNotNone(db.get(InstalacionCamara,1).fecha_registro)
        page=self.web.get(r.json()['url']);self.assertEqual(page.status_code,200,page.text)
        for word in ['Cable UTP','12.50','TP-Link C200','Ana Técnica']:self.assertIn(word,page.text)
        self.assertEqual(self.web.get('/camaras/reportes?desde=2026-11-01').status_code,200)
        self.assertIn('0 instalaciones',self.web.get('/camaras/reportes?desde=2026-11-01').text)
        self.assertEqual(self.web.get('/camaras/reportes?desde=2026-12-01&hasta=2026-01-01').status_code,422)
        self.assertEqual(self.web.get('/camaras/999/reporte').status_code,404)
    def test_validacion_y_sin_productos(self):
        for patch in [dict(tecnico_id=999),dict(celular='123'),dict(productos=[dict(descripcion='Cable',cantidad='-1',unidad='Metro')]),dict(fecha_registro='ayer')]:
            self.assertEqual(self.web.post('/camaras/api',json={**self.data,**patch}).status_code,422)
        r=self.web.post('/camaras/api',json={**self.data,'productos':[]});self.assertEqual(r.status_code,201,r.text)
        self.assertIn('No se registraron productos',self.web.get(r.json()['url']).text)
        self.assertEqual(self.web.post('/camaras/api',json={**self.data,'direccion':'Otra direccion'}).status_code,409)
    def test_clientes_historial_instalaciones(self):
        for numero in range(11):
            r=self.web.post('/camaras/api',json={**self.data,'solicitud_id':str(uuid4())})
            self.assertEqual(r.status_code,201,r.text)
        r=self.web.get('/clientes/api/instalaciones?q=Ana')
        self.assertEqual(r.status_code,200,r.text)
        datos=r.json();self.assertEqual(datos['total'],1)
        cliente=datos['clientes'][0];self.assertEqual(cliente['total_reparaciones'],11)
        self.assertEqual(self.web.get('/clientes/api/instalaciones?q=NoExiste').json()['total'],0)
        self.assertEqual(self.web.get('/clientes/api/instalaciones?q=12345678').json()['total'],1)
        url=f"/clientes/api/instalaciones/{cliente['id']}/historial"
        h=self.web.get(url).json();self.assertEqual(h['total'],11);self.assertEqual(len(h['instalaciones']),10)
        self.assertEqual(len(self.web.get(url+'?pagina=2').json()['instalaciones']),1)
        self.assertEqual(self.web.get(h['instalaciones'][0]['url']).status_code,200)
        self.assertEqual(self.web.get('/clientes/api/instalaciones/999/historial').status_code,404)
        self.assertEqual(self.web.get('/clientes/api').status_code,200)

    def test_otro_servicio_y_domicilio(self):
        r=self.web.post('/camaras/api',json={**self.data,'productos':[], 'otro_producto_servicio':'Configurar red', 'servicio_domicilio':'si'})
        self.assertEqual(r.status_code,201,r.text)
        with fixtures.SessionLocal() as db:
            self.assertEqual(db.get(InstalacionCamara,1).datos['servicio_domicilio'],'si')
        self.assertIn('Configurar red',self.web.get(r.json()['url']).text)
        self.assertEqual(self.web.get(r.json()['url']+'.pdf').status_code,200)

    def test_fecha_pasada_rechazada(self):
        from datetime import datetime, timedelta
        from app.routes.camaras import PERU
        ayer=(datetime.now(PERU).date()-timedelta(days=1)).isoformat()
        r=self.web.post('/camaras/api',json={**self.data,'fecha_instalacion':ayer})
        self.assertEqual(r.status_code,422,r.text)

    def test_camaras_opcionales(self):
        for cams in [[],[{'modelo':'Solo modelo'}],[{'cantidad':2}]]:
            r=self.web.post('/camaras/api',json={**self.data,'solicitud_id':str(uuid4()),'camaras':cams})
            self.assertEqual(r.status_code,201,r.text)
            self.assertEqual(self.web.get(r.json()['url']+'.pdf').status_code,200)
        self.assertEqual(self.web.post('/camaras/api',json={**self.data,'camaras':[{'cantidad':0}]}).status_code,422)

    def test_excel_texto_seguro(self):
        self.data['observaciones']='=1+1';self.web.post('/camaras/api',json=self.data)
        r=self.web.get('/camaras/reportes.xlsx');self.assertEqual(r.status_code,200)
        wb=load_workbook(BytesIO(r.content));self.assertEqual(wb.sheetnames,['Instalaciones','Cámaras','Productos vendidos'])
        self.assertEqual(wb['Instalaciones']['J2'].value,'=1+1');self.assertEqual(wb['Instalaciones']['J2'].data_type,'s')
        self.assertEqual(wb['Productos vendidos']['C2'].value,12.5)
        self.assertIn('Ana Técnica',self.web.get('/camaras').text)

if __name__=='__main__':unittest.main()
