from datetime import date, datetime, timezone, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Literal
from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import Response
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel, ConfigDict, Field, field_validator
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.database import obtener_db
from app.models.instalacion import InstalacionCamara
from app.models.trabajador import Trabajador

router = APIRouter(prefix='/camaras', tags=['Instalación de cámaras'])
templates = Jinja2Templates(directory=Path(__file__).resolve().parent.parent / 'templates')
PERU = timezone(timedelta(hours=-5))


def fecha_local(valor):
    if valor.tzinfo is None:
        valor = valor.replace(tzinfo=timezone.utc)
    return valor.astimezone(PERU).strftime('%d/%m/%Y %H:%M')


templates.env.filters['fecha_instalacion_local'] = fecha_local


class Texto(BaseModel):
    model_config = ConfigDict(extra='forbid', str_strip_whitespace=True)


class Camara(Texto):
    modelo: str = Field(default='', max_length=150)
    procedencia: Literal['ServiTech', 'Cliente', ''] = ''
    cantidad: int | None = Field(default=None, ge=1, le=1000, strict=True)


class Producto(Texto):
    descripcion: str = Field(min_length=1, max_length=150)
    cantidad: Decimal = Field(gt=0, le=100000, max_digits=9, decimal_places=2)
    unidad: Literal['Unidad', 'Metro', 'Caja', 'Rollo', 'Otro']


class Registro(Texto):
    solicitud_id: UUID
    nombres: str = Field(min_length=1, max_length=100)
    apellidos: str = Field(min_length=1, max_length=100)
    dni_ruc: str = Field(pattern=r'^(?:[0-9]{8}|[0-9]{11})$')
    celular: str = Field(pattern=r'^[0-9]{9}$')
    tecnico_id: int = Field(gt=0)
    fecha_instalacion: date
    direccion: str = Field(min_length=5, max_length=500)
    camaras: list[Camara] = Field(default_factory=list, max_length=100)
    observaciones: str = Field(default='', max_length=3000)
    otro_producto_servicio: str = Field(default='', max_length=3000)
    servicio_domicilio: Literal['si', 'no'] = 'no'
    productos: list[Producto] = Field(default_factory=list, max_length=100)


def filtrar(db, desde, hasta, tecnico):
    if desde and hasta and desde > hasta:
        raise HTTPException(422, 'La fecha inicial no puede ser posterior a la final.')
    q = db.query(InstalacionCamara)
    if desde: q = q.filter(InstalacionCamara.fecha_instalacion >= desde)
    if hasta: q = q.filter(InstalacionCamara.fecha_instalacion <= hasta)
    if tecnico: q = q.filter(InstalacionCamara.tecnico == tecnico)
    return q.order_by(InstalacionCamara.fecha_instalacion.desc(), InstalacionCamara.id.desc())


@router.get('')
def ver_camaras(request: Request, db: Session = Depends(obtener_db)):
    trabajadores = db.query(Trabajador).filter(Trabajador.activo.is_(True)).order_by(Trabajador.nombre).all()
    return templates.TemplateResponse(request=request, name='camaras.html', context={'seccion':'camaras','trabajadores':trabajadores,'hoy':datetime.now(PERU).date().isoformat()})


@router.post('/api', status_code=201)
def registrar_instalacion(datos: Registro, db: Session = Depends(obtener_db)):
    contenido = datos.model_dump(mode='json', exclude={'solicitud_id'})
    token = str(datos.solicitud_id)
    existente = db.query(InstalacionCamara).filter_by(solicitud_id=token).first()
    if existente:
        if existente.datos != contenido: raise HTTPException(409, 'Este envío ya fue utilizado. Recarga el formulario antes de crear otro registro.')
        return {'numero':existente.numero,'url':f'/camaras/{existente.id}/reporte'}
    if datos.fecha_instalacion < datetime.now(PERU).date():
        raise HTTPException(422, 'La fecha programada no puede ser anterior a hoy (hora de Perú).')
    tecnico = db.query(Trabajador).filter_by(id=datos.tecnico_id, activo=True).first()
    if tecnico is None: raise HTTPException(422, 'Selecciona un técnico activo de Ajustes.')
    registro = InstalacionCamara(solicitud_id=token, fecha_instalacion=datos.fecha_instalacion, tecnico=tecnico.nombre, datos=contenido)
    db.add(registro)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        existente=db.query(InstalacionCamara).filter_by(solicitud_id=token).first()
        if existente and existente.datos == contenido:
            return {'numero':existente.numero,'url':f'/camaras/{existente.id}/reporte'}
        raise HTTPException(409, 'No se pudo guardar el registro. Revisa si ya fue registrado.') from None
    return {'numero':registro.numero,'url':f'/camaras/{registro.id}/reporte'}


@router.get('/reportes')
def reportes(request: Request, desde: date | None = None, hasta: date | None = None, tecnico: str = '', pagina: int = Query(1, ge=1), db: Session = Depends(obtener_db)):
    q=filtrar(db,desde,hasta,tecnico); total=q.count(); paginas=max(1,(total+19)//20); pagina=min(pagina,paginas)
    tecnicos=[x[0] for x in db.query(InstalacionCamara.tecnico).distinct().order_by(InstalacionCamara.tecnico).all()]
    return templates.TemplateResponse(request=request,name='camaras_reportes.html',context=dict(seccion='camaras',registros=q.offset((pagina-1)*20).limit(20).all(),total=total,pagina=pagina,paginas=paginas,desde=desde,hasta=hasta,tecnico=tecnico,tecnicos=tecnicos))


@router.get('/reportes.xlsx')
def exportar(desde: date | None = None, hasta: date | None = None, tecnico: str = '', db: Session = Depends(obtener_db)):
    registros=filtrar(db,desde,hasta,tecnico).all()
    wb=Workbook();ws=wb.active;ws.title='Instalaciones'
    ws.append(['Registro','Fecha de registro (Perú)','Fecha programada','Nombres','Apellidos','DNI/RUC','Celular','Técnico','Dirección','Observaciones'])
    cam=wb.create_sheet('Cámaras');cam.append(['Registro','Modelo','Procedencia','Cantidad'])
    prod=wb.create_sheet('Productos vendidos');prod.append(['Registro','Producto','Cantidad','Unidad'])
    def fila(hoja,valores):
        hoja.append(valores)
        for cell in hoja[hoja.max_row]:
            if isinstance(cell.value,str): cell.data_type='s'
    for r in registros:
        d=r.datos
        fila(ws,[r.numero,fecha_local(r.fecha_registro),r.fecha_instalacion.isoformat(),d['nombres'],d['apellidos'],d['dni_ruc'],d['celular'],r.tecnico,d['direccion'],d['observaciones']])
        for x in d['camaras']: fila(cam,[r.numero,x['modelo'],x['procedencia'],x['cantidad']])
        for x in d['productos']: fila(prod,[r.numero,x['descripcion'],float(x['cantidad']),x['unidad']])
    for hoja in wb:
        hoja.freeze_panes='A2';hoja.auto_filter.ref=hoja.dimensions
        for cell in hoja[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='CE0033')
        for col in hoja.columns:hoja.column_dimensions[col[0].column_letter].width=min(45,max(16,max(len(str(c.value or '')) for c in col)+2))
    out=BytesIO();wb.save(out)
    return Response(out.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="Instalaciones-camaras.xlsx"','Cache-Control':'no-store'})


@router.get('/{registro_id}/reporte')
def reporte_individual(registro_id: int, request: Request, db: Session = Depends(obtener_db)):
    registro=db.get(InstalacionCamara,registro_id)
    if registro is None:raise HTTPException(404,'Instalación no encontrada.')
    return templates.TemplateResponse(request=request,name='camara_reporte.html',context={'seccion':'camaras','registro':registro})


@router.get('/{registro_id}/reporte.pdf')
def descargar_reporte(registro_id: int, db: Session = Depends(obtener_db)):
    registro = db.get(InstalacionCamara, registro_id)
    if registro is None:
        raise HTTPException(404, 'Instalación no encontrada.')
    from app.services.instalacion_pdf import generar_instalacion_pdf
    return Response(generar_instalacion_pdf(registro, fecha_local(registro.fecha_registro)), media_type='application/pdf', headers={'Content-Disposition': f'attachment; filename="Instalacion-{registro.numero}.pdf"', 'Cache-Control': 'no-store'})
