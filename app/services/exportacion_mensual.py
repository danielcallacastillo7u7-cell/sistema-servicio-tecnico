from datetime import datetime, timezone, timedelta
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from app.models.instalacion import InstalacionCamara
from app.models.orden import OrdenServicio
from app.models.trabajador import Trabajador
from app.services.exportaciones import obtener_filas
import json

MESES=['Enero','Febrero','Marzo','Abril','Mayo','Junio','Julio','Agosto','Septiembre','Octubre','Noviembre','Diciembre']
TIPOS={'ordenes':'Órdenes de servicio','instalaciones':'Instalaciones de cámaras','ambos':'Órdenes e instalaciones'}
def tecnicos_disponibles(db):
    nombres={x[0] for x in db.query(Trabajador.nombre).all()}
    nombres.update(x[0] for x in db.query(OrdenServicio.tecnico_responsable).distinct().all())
    nombres.update(x[0] for x in db.query(InstalacionCamara.tecnico).distinct().all())
    return sorted(n for n in nombres if n)

def crear_mensual(db,tipo,tecnicos):
    if tipo not in TIPOS:raise ValueError('Selecciona órdenes, instalaciones o ambos.')
    seleccion=list(dict.fromkeys(t.strip() for t in tecnicos if t.strip()))
    if len(seleccion)>2:raise ValueError('Selecciona como máximo dos técnicos.')
    if any(t not in tecnicos_disponibles(db) for t in seleccion):raise ValueError('Selecciona técnicos registrados en el sistema.')
    grupos={};cantidad=0
    def agregar(clase,fecha,encabezado,fila):
        key=(fecha.year,fecha.month,clase)
        if key not in grupos:grupos[key]=[encabezado]
        grupos[key].append(fila)
    if tipo in ('ordenes','ambos'):
        for f in obtener_filas(db,'completo'):
            if seleccion and f['tecnico_responsable'] not in seleccion:continue
            fecha=datetime.fromisoformat(f['fecha_ingreso'])
            if fecha.tzinfo is None:fecha=fecha.replace(tzinfo=timezone.utc)
            fecha=fecha.astimezone(timezone(timedelta(hours=-5)))
            equipos=json.loads(f['equipos_recibidos'])
            detalle='\n'.join(' | '.join(str(e.get(k) or '') for k in ['tipo','marca','modelo','numero_serie','accesorios','observaciones']) for e in equipos)
            agregar('Órdenes',fecha,['Orden','Fecha de ingreso (Perú)','Cliente','DNI/RUC','Celular','Técnico','Estado','Falla reportada','Equipos y accesorios','Falla encontrada','Solución recomendada','Costo estimado'],[f['numero_orden'],fecha.strftime('%d/%m/%Y %H:%M'),f['cliente'],f['dni_ruc'],f['telefono'],f['tecnico_responsable'],f['estado'],f['falla_reportada'],detalle,f['falla_encontrada'],f['solucion_recomendada'],f['costo_estimado']]);cantidad+=1
    if tipo in ('instalaciones','ambos'):
        q=db.query(InstalacionCamara)
        if seleccion:q=q.filter(InstalacionCamara.tecnico.in_(seleccion))
        for reg in q.order_by(InstalacionCamara.fecha_instalacion,InstalacionCamara.id).all():
            d=reg.datos;fecha=reg.fecha_registro
            if fecha.tzinfo is None:fecha=fecha.replace(tzinfo=timezone.utc)
            fecha=fecha.astimezone(timezone(timedelta(hours=-5)))
            cams='\n'.join(f"{x['modelo']} | {x['procedencia']} | Cantidad: {x['cantidad']}" for x in d['camaras'])
            productos='\n'.join(f"{x['descripcion']} | {x['cantidad']} {x['unidad']}" for x in d['productos'])
            agregar('Instalaciones',reg.fecha_instalacion,['Registro','Fecha programada','Fecha de registro (Perú)','Cliente','DNI/RUC','Celular','Técnico','Dirección','Cámaras','Observaciones','Productos vendidos'],[reg.numero,reg.fecha_instalacion.strftime('%d/%m/%Y'),fecha.strftime('%d/%m/%Y %H:%M'),d['nombres']+' '+d['apellidos'],d['dni_ruc'],d['celular'],reg.tecnico,d['direccion'],cams,d['observaciones'],productos]);cantidad+=1
    wb=Workbook();wb.remove(wb.active)
    for (year,month,clase),filas in sorted(grupos.items()):
        ws=wb.create_sheet(f'{clase} {MESES[month-1]} {year}')
        for fila in filas:
            ws.append(fila)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str):cell.data_type='s'
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17212D')
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(45,max(18,max(len(str(c.value or '')) for c in col)+2))
    if not grupos:wb.create_sheet('Sin registros').append(['No hay registros para la selección de tipo y técnicos.'])
    out=BytesIO();wb.save(out);out.seek(0)
    return out,cantidad,seleccion
